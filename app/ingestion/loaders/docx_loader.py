"""
docx_loader.py

Loads Microsoft Word documents.
"""

from pathlib import Path

from docx import Document as DocxDocument

from app.ingestion.loaders.base_loader import BaseLoader

from app.schemas.document import Document
from app.schemas.page import Page
from app.schemas.metadata import Metadata


class DOCXLoader(BaseLoader):

    def load(
        self,
        file_path: Path,
    ) -> Document:


        docx = DocxDocument(
            str(file_path)
        )


        text = "\n".join(
            paragraph.text
            for paragraph in docx.paragraphs
            if paragraph.text.strip()
        )


        return Document(
            filename=file_path.name,
            filepath=str(file_path),

            metadata=Metadata(),

            pages=[
                Page(
                    page_number=1,
                    text=text,
                )
            ],
        )