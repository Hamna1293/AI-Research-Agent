from pathlib import Path
from typing import cast

from pptx import Presentation
from pptx.shapes.autoshape import Shape

from app.ingestion.loaders.base_loader import BaseLoader

from app.schemas.document import Document
from app.schemas.page import Page
from app.schemas.metadata import Metadata


class PPTLoader(BaseLoader):

    def load(
        self,
        file_path: Path,
    ) -> Document:


        presentation = Presentation(
            str(file_path)
        )


        slides = []


        for index, slide in enumerate(
            presentation.slides,
            start=1
        ):

            slide_text = []


            for shape in slide.shapes:

                if shape.has_text_frame:

                    text_shape = cast(
                        Shape,
                        shape
                    )

                    slide_text.append(
                        text_shape.text
                    )


            slides.append(
                f"Slide {index}\n"
                +
                "\n".join(slide_text)
            )


        text = "\n\n".join(
            slides
        )


        return Document(
            filename=file_path.name,
            filepath=str(file_path),

            metadata=Metadata(),

            pages=[
                Page(
                    page_number=1,
                    text=text,
                )
            ],
        )